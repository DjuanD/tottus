import os
import telebot
from telebot import types
import pandas as pd
from flask import Flask
import threading
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
import time
from openpyxl.styles import PatternFill
from openpyxl import load_workbook

TOKEN = os.getenv("TELEGRAM_TOKEN")
bot = telebot.TeleBot(TOKEN)
user_step = {}

# Cargar productos y tiendas
productos_df = pd.read_excel("productos.xlsx")
tiendas_df = pd.read_excel("tiendas_tottus.xlsx")

# Inicia scraping
def hacer_scraping(filtro, tipo):
    resultados = []

    productos_filtrados = productos_df[productos_df[tipo].str.contains(filtro, case=False, na=False)]
    for _, producto in productos_filtrados.iterrows():
        codigo = str(producto['CODIGO'])
        nombre_producto = producto['PRODUCTO']
        marca = producto['MARCA']

        for _, tienda in tiendas_df.iterrows():
            url = tienda['tienda']
            nombre_tienda = tienda['manual']

            options = Options()
            options.add_argument("--headless")
            options.add_argument("--disable-gpu")
            options.add_argument("--no-sandbox")
            options.add_argument("--window-size=375,812")  # Simula móvil

            driver = webdriver.Chrome(options=options)
            try:
                driver.get(url)
                search_input = driver.find_element(By.NAME, "search")  # Ajusta si cambia
                search_input.send_keys(codigo)
                search_input.submit()

                time.sleep(10)  # Esperar carga

                precio_element = driver.find_element(By.CSS_SELECTOR, ".product-price")
                precio = precio_element.text
            except Exception:
                precio = "No encontrado"
            finally:
                driver.quit()

            resultados.append({
                'CODIGO': codigo,
                'PRODUCTO': nombre_producto,
                'MARCA': marca,
                'TIENDA': nombre_tienda,
                'PRECIO': precio
            })

    df_resultado = pd.DataFrame(resultados)
    archivo = "resultados.xlsx"
    df_resultado.to_excel(archivo, index=False)

    # Formato condicional
    wb = load_workbook(archivo)
    ws = wb.active
    verde = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    amarillo = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")

    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            if "S/" in str(cell.value):
                precio = float(cell.value.replace("S/", "").replace(",", ".").strip())
                if precio < 10:
                    cell.fill = verde
                elif precio < 20:
                    cell.fill = amarillo

    wb.save(archivo)
    return archivo

# ---- Telegram ----

@bot.message_handler(commands=['start'])
def start(message):
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True, one_time_keyboard=True)
    markup.add("Buscar por producto", "Buscar por marca")
    bot.send_message(message.chat.id, "¿Cómo quieres buscar?", reply_markup=markup)

@bot.message_handler(func=lambda m: m.text in ["Buscar por producto", "Buscar por marca"])
def elegir_busqueda(message):
    tipo = "PRODUCTO" if "producto" in message.text.lower() else "MARCA"
    user_step[message.chat.id] = tipo
    bot.send_message(message.chat.id, f"Escribe el {tipo.lower()} que deseas buscar:")

@bot.message_handler(func=lambda m: m.chat.id in user_step)
def recibir_filtro(message):
    tipo = user_step.pop(message.chat.id)
    bot.send_message(message.chat.id, "Buscando, espera unos segundos...")

    archivo = hacer_scraping(message.text, tipo)
    with open(archivo, "rb") as file:
        bot.send_document(message.chat.id, file)

# ---- Flask y threading para Cloud Run ----

app = Flask(__name__)

@app.route('/')
def home():
    return "Bot Tottus activo en Cloud Run"

def iniciar_bot():
    bot.infinity_polling()

if __name__ == '__main__':
    threading.Thread(target=iniciar_bot).start()
    app.run(host='0.0.0.0', port=8080)
